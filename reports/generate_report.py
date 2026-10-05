#!/usr/bin/env python3
"""Pulse weekly progress report renderer.

Laravel builds the weekly snapshot (app/Services/PortfolioSnapshot.php), writes it
to a JSON file and calls:

    python3 generate_report.py --input snapshot.json --output report.pdf  --format pdf
    python3 generate_report.py --input snapshot.json --output report.pptx --format pptx

Requires: reportlab, python-pptx  (pip install -r requirements.txt)
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import date

RAG_LABEL = {"G": "On track", "A": "At risk", "R": "Off track", None: "No update"}
RAG_HEX = {"G": "#1E8E5A", "A": "#D08A00", "R": "#C8372D", None: "#8592A0"}
RAG_SOFT = {"G": "#E3F4EB", "A": "#FFF3DC", "R": "#FDE8E6", None: "#EEF1F4"}
SEV_HEX = {"Critical": "#C8372D", "High": "#E0663A", "Medium": "#D08A00", "Low": "#8592A0"}
INK, MUTED, LINE = "#16212B", "#5D6B78", "#E1E6EB"
TREND = {"up": "improved", "down": "worsened", "flat": "steady", "new": "new", "none": ""}


def fmt_date(s: str | None) -> str:
    if not s:
        return "—"
    try:
        return date.fromisoformat(s[:10]).strftime("%d %b")
    except ValueError:
        return s


def fmt_num(v) -> str:
    if v is None:
        return "—"
    v = float(v)
    return str(int(v)) if v.is_integer() else f"{v:.1f}"


def metric_ratio(m) -> float | None:
    tgt = m.get("target") or (100 if (m.get("unit") or "") == "%" else None)
    if not tgt:
        return None
    return max(0.0, min(1.0, float(m["value"]) / float(tgt)))


def metric_color(r: float | None) -> str:
    if r is None:
        return "#0B4F6C"
    return RAG_HEX["G"] if r >= 0.85 else RAG_HEX["A"] if r >= 0.5 else RAG_HEX["R"]


def scope_label(snap) -> str:
    return snap.get("programme") or "All projects"


# =============================================================================== PDF
def render_pdf(snap: dict, out: str) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)
    from reportlab.platypus.flowables import Flowable

    brand = colors.HexColor(snap.get("brand") or "#0B4F6C")
    C = colors.HexColor
    page = landscape(A4)
    W = page[0] - 28 * mm

    def st(name, **kw):
        base = dict(fontName="Helvetica", fontSize=9, leading=12, textColor=C(INK), alignment=TA_LEFT)
        base.update(kw)
        return ParagraphStyle(name, **base)

    S = {
        "h1": st("h1", fontName="Helvetica-Bold", fontSize=15, leading=19),
        "h2": st("h2", fontName="Helvetica-Bold", fontSize=11.5, leading=15, spaceBefore=4, spaceAfter=4, keepWithNext=1),
        "body": st("body"),
        "small": st("small", fontSize=8, leading=10.5, textColor=C(MUTED)),
        "cell": st("cell", fontSize=8.5, leading=11),
        "cellb": st("cellb", fontName="Helvetica-Bold", fontSize=8.5, leading=11),
        "th": st("th", fontName="Helvetica-Bold", fontSize=8, leading=10, textColor=C(MUTED)),
        "lbl": st("lbl", fontName="Helvetica-Bold", fontSize=7.5, leading=10, textColor=C(MUTED)),
        "kv": st("kv", fontName="Helvetica-Bold", fontSize=18, leading=21),
        "bul": st("bul", fontSize=8.5, leading=11.5, leftIndent=9, bulletIndent=0),
    }

    def esc(s):
        return (str(s or "")).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    class Chip(Flowable):
        """Rounded RAG chip."""
        def __init__(self, rag, w=62, h=13):
            super().__init__()
            self.rag, self.width, self.height = rag, w, h

        def draw(self):
            c = self.canv
            c.setFillColor(C(RAG_SOFT[self.rag]))
            c.roundRect(0, 0, self.width, self.height, 6.5, stroke=0, fill=1)
            c.setFillColor(C(RAG_HEX[self.rag]))
            c.circle(8, self.height / 2, 3, stroke=0, fill=1)
            c.setFont("Helvetica-Bold", 7.5)
            c.drawString(15, 3.6, RAG_LABEL[self.rag])

    class Bar(Flowable):
        def __init__(self, ratio, w=60, h=5, color=None):
            super().__init__()
            self.ratio, self.width, self.height, self.color = ratio or 0, w, h, color or brand

        def draw(self):
            c = self.canv
            c.setFillColor(C("#E8ECF0"))
            c.roundRect(0, 0, self.width, self.height, 2.5, stroke=0, fill=1)
            if self.ratio > 0:
                c.setFillColor(self.color)
                c.roundRect(0, 0, max(self.height, self.width * self.ratio), self.height, 2.5, stroke=0, fill=1)

    def bullets(items, color=None):
        if not items:
            return [Paragraph("—", S["small"])]
        style = ParagraphStyle("bul_" + (color or MUTED), parent=S["bul"], bulletColor=C(color or MUTED),
                               bulletFontName="Helvetica-Bold")
        return [Paragraph(esc(i), style, bulletText="•") for i in items]

    def trend_html(t):
        col = {"up": RAG_HEX["G"], "down": RAG_HEX["R"]}.get(t, MUTED)
        return f'<font color="{col}"><b>{TREND.get(t, "")}</b></font>' if t in ("up", "down") else TREND.get(t, "")

    def header_footer(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(brand)
        canvas.rect(0, page[1] - 9 * mm, page[0], 9 * mm, stroke=0, fill=1)
        canvas.setFillColor(colors.white)
        canvas.setFont("Helvetica-Bold", 9)
        canvas.drawString(14 * mm, page[1] - 6 * mm, f"{snap.get('org', '')}  ·  {snap.get('title', '')}")
        canvas.setFont("Helvetica", 9)
        canvas.drawRightString(page[0] - 14 * mm, page[1] - 6 * mm, f"Week of {snap['week_label']}  ·  {scope_label(snap)}")
        canvas.setFillColor(C(MUTED))
        canvas.setFont("Helvetica", 7.5)
        canvas.drawString(14 * mm, 8 * mm, f"Generated {snap.get('generated_at', '')} from Pulse")
        canvas.drawRightString(page[0] - 14 * mm, 8 * mm, f"Page {doc.page}")
        canvas.restoreState()

    t = snap["totals"]
    projects = snap["projects"]
    story = [Paragraph(f"Weekly progress — {esc(snap['week_label'])}", S["h1"]),
             Paragraph(f"{esc(scope_label(snap))} · {t['projects']} projects tracked", S["small"]), Spacer(1, 8)]

    # ---- KPI tiles
    tiles = [("ON TRACK", t["G"], RAG_HEX["G"]), ("AT RISK", t["A"], RAG_HEX["A"]), ("OFF TRACK", t["R"], RAG_HEX["R"]),
             ("UPDATES IN", f"{t['submitted']}/{t['expected']}", INK),
             ("OPEN ISSUES", t["open_issues"], INK), ("HIGH / CRITICAL", t["critical_issues"], RAG_HEX["R"] if t["critical_issues"] else INK),
             ("OVERDUE MILESTONES", t["overdue_milestones"], RAG_HEX["A"] if t["overdue_milestones"] else INK)]
    cells = [[Paragraph(lbl, S["lbl"]), Paragraph(f'<font color="{col}">{esc(val)}</font>', S["kv"])] for lbl, val, col in tiles]
    kt = Table([cells], colWidths=[W / len(tiles)] * len(tiles))
    kt.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.6, C(LINE)), ("INNERGRID", (0, 0), (-1, -1), 0.6, C(LINE)),
                            ("BACKGROUND", (0, 0), (-1, -1), C("#FAFBFC")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    story += [kt, Spacer(1, 10)]

    # ---- portfolio table
    story.append(Paragraph("Portfolio at a glance", S["h2"]))
    rows = [[Paragraph(h, S["th"]) for h in ("Project", "Phase", "Status", "Trend", "Progress", "Owner", "Headline")]]
    for p in projects:
        prog = [Paragraph(f"{p['progress']}%" if p.get("progress") is not None else "—", S["cell"])]
        if p.get("progress") is not None:
            prog.append(Bar(p["progress"] / 100, w=42, color=brand))
        name = f"<b>{esc(p['name'])}</b>"
        if p.get("stale"):
            name += f'<br/><font color="{RAG_HEX["A"]}" size="7">No update this week — showing {fmt_date(p.get("update_week")) if p.get("update_week") else "nothing yet"}</font>'
        rows.append([Paragraph(name, S["cell"]), Paragraph(esc(p.get("phase")), S["cell"]), Chip(p.get("rag")),
                     Paragraph(trend_html(p.get("trend")), S["small"]), prog,
                     Paragraph(esc(p.get("owner") or "—"), S["cell"]), Paragraph(esc(p.get("summary") or "—"), S["cell"])])
    pt = Table(rows, colWidths=[W * .22, W * .07, W * .09, W * .07, W * .08, W * .11, W * .36], repeatRows=1)
    pt.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.5, C(LINE)), ("BACKGROUND", (0, 0), (-1, 0), C("#F4F6F8")),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 5),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    story += [pt, Spacer(1, 10)]

    # ---- issues table
    def issue_table(title, items):
        if not items:
            return []
        rows = [[Paragraph(h, S["th"]) for h in ("Issue", "Project", "Severity", "Status", "Owner", "Due", "Latest progress")]]
        for i in items:
            due = fmt_date(i.get("due_date"))
            if i.get("overdue"):
                due = f'<font color="{RAG_HEX["R"]}"><b>{due} overdue</b></font>'
            rows.append([Paragraph(f"<b>{esc(i['title'])}</b><br/><font size='7' color='{MUTED}'>{esc(i['kind'])}</font>", S["cell"]),
                         Paragraph(esc(i.get("project") or "—"), S["cell"]),
                         Paragraph(f'<font color="{SEV_HEX.get(i["severity"], MUTED)}"><b>{esc(i["severity"])}</b></font>', S["cell"]),
                         Paragraph(esc(i["status"]), S["cell"]), Paragraph(esc(i.get("owner") or "—"), S["cell"]),
                         Paragraph(due, S["cell"]), Paragraph(esc(i.get("latest_note") or "—"), S["cell"])])
        tb = Table(rows, colWidths=[W * .22, W * .14, W * .07, W * .08, W * .12, W * .08, W * .29], repeatRows=1)
        tb.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.5, C(LINE)), ("BACKGROUND", (0, 0), (-1, 0), C("#F4F6F8")),
                                ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 5),
                                ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
        return [Paragraph(title, S["h2"]), tb, Spacer(1, 10)]

    prod = [i for i in snap["issues"] if i["kind"] == "Production issue"]
    raid = [i for i in snap["issues"] if i["kind"] != "Production issue"]
    story += issue_table("Production issues being monitored", prod)
    story += issue_table("Open risks, issues & dependencies", raid)
    if snap.get("resolved"):
        story.append(Paragraph("Resolved this week: " + "; ".join(esc(i["title"]) for i in snap["resolved"]), S["small"]))

    # ---- project detail pages
    story.append(PageBreak())
    story.append(Paragraph("Project detail", S["h1"]))  # projects ordered: off track first
    story.append(Spacer(1, 6))
    for p in projects:
        head = Table([[Paragraph(f"<b>{esc(p['name'])}</b>"
                                 f"<font color='{MUTED}' size='8'>   {esc(p.get('code') or '')} · {esc(p.get('phase'))} · Owner: {esc(p.get('owner') or '—')}"
                                 f"{' · Target ' + fmt_date(p.get('target_date')) if p.get('target_date') else ''}</font>", st("ph", fontSize=11, leading=14)),
                       Chip(p.get("rag"), w=66)]], colWidths=[W - 72, 72])
        head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LINEBELOW", (0, 0), (-1, 0), 1.2, C(RAG_HEX[p.get("rag")])),
                                  ("BOTTOMPADDING", (0, 0), (-1, -1), 5), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
        block = [head, Spacer(1, 4)]
        if p.get("stale"):
            block.append(Paragraph(f'<font color="{RAG_HEX["A"]}">No update submitted for this week — figures below are from '
                                   f'{fmt_date(p.get("update_week")) if p.get("update_week") else "no previous update"}.</font>', S["small"]))
        if p.get("summary"):
            block.append(Paragraph(esc(p["summary"]), S["body"]))
        block.append(Spacer(1, 5))

        # metrics column
        mrows = []
        for m in p.get("metrics") or []:
            r = metric_ratio(m)
            tgt = f" / {fmt_num(m['target'])}{m.get('unit') or ''}" if m.get("target") not in (None, 100) or (m.get("unit") or "") != "%" else ""
            mrows.append([Paragraph(esc(m["label"]), S["cell"]),
                          Paragraph(f"<b>{fmt_num(m['value'])}{esc(m.get('unit') or '')}</b><font color='{MUTED}' size='7'>{esc(tgt)}</font>", S["cell"])])
            if r is not None:
                mrows.append([Bar(r, w=W * .26 - 12, color=C(metric_color(r))), ""])
        left = [Paragraph("KEY METRICS", S["lbl"])]
        if mrows:
            mt = Table(mrows, colWidths=[W * .19, W * .08])
            mt.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 1),
                                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
            left.append(mt)
        else:
            left.append(Paragraph("—", S["small"]))
        ms = p.get("milestones_overdue", []) + [m for m in p.get("milestones_due", []) if m not in p.get("milestones_overdue", [])]
        left += [Spacer(1, 6), Paragraph("MILESTONES — OVERDUE &amp; NEXT 3 WEEKS", S["lbl"])]
        if ms:
            for m in ms[:6]:
                late = m in p.get("milestones_overdue", [])
                flag = '<font color="%s"><b>Overdue</b></font> ' % RAG_HEX["R"] if late else ""
                left.append(Paragraph("%s%s <font color='%s'>· %s · %s</font>" % (
                    flag, esc(m["title"]), MUTED, fmt_date(m["due_date"]), esc(m["status"])), S["bul"], bulletText="•"))
        else:
            left.append(Paragraph("—", S["small"]))
        left.append(Paragraph(f"{p.get('milestones_done', 0)} of {p.get('milestones_total', 0)} milestones complete", S["small"]))

        mid = [Paragraph("KEY ITEMS IN THE CRITICAL PATH", S["lbl"])] + bullets(p.get("critical_path"), RAG_HEX["R"])
        if p.get("support_needed"):
            mid += [Spacer(1, 6), Paragraph("SUPPORT NEEDED", S["lbl"])] + bullets(p["support_needed"], RAG_HEX["A"])
        right = ([Paragraph("ACHIEVED THIS WEEK", S["lbl"])] + bullets(p.get("achievements"), RAG_HEX["G"])
                 + [Spacer(1, 6), Paragraph("NEXT STEPS", S["lbl"])] + bullets(p.get("next_steps"), MUTED))

        body = Table([[left, mid, right]], colWidths=[W * .30, W * .35, W * .35])
        body.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (0, -1), 0),
                                  ("LINEBEFORE", (1, 0), (-1, -1), 0.5, C(LINE)), ("LEFTPADDING", (1, 0), (-1, -1), 10)]))
        block += [body, Spacer(1, 14)]
        story.append(KeepTogether(block))

    doc = SimpleDocTemplate(out, pagesize=page, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=15 * mm,
                            bottomMargin=14 * mm, title=f"{snap.get('title')} — {snap['week_label']}", author=snap.get("org", ""))
    doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)


# =============================================================================== PPTX
def render_pptx(snap: dict, out: str) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    def rgb(h):
        h = h.lstrip("#")
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    brand = snap.get("brand") or "#0B4F6C"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    SW = prs.slide_width

    def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
        s = slide.shapes.add_shape(shape, x, y, w, h)
        if fill:
            s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
        else:
            s.fill.background()
        if line:
            s.line.color.rgb = rgb(line); s.line.width = Pt(0.75)
        else:
            s.line.fill.background()
        s.shadow.inherit = False
        return s

    def text(slide, x, y, w, h, value, size=12, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        tf.vertical_anchor = anchor
        p = tf.paragraphs[0]
        p.alignment = align
        r = p.add_run(); r.text = str(value)
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, rgb(color), "Calibri"
        return tb

    def bullet_list(slide, x, y, w, h, title, items, dot_color, size=12, max_items=7):
        text(slide, x, y, w, Inches(0.3), title.upper(), 10, True, MUTED)
        tb = slide.shapes.add_textbox(x, y + Inches(0.3), w, h - Inches(0.3))
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = Inches(0.05)
        items = items or ["—"]
        shown = items[:max_items] + ([f"+{len(items) - max_items} more in Pulse"] if len(items) > max_items else [])
        for i, it in enumerate(shown):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(5)
            if it != "—":
                d = p.add_run(); d.text = "●  "; d.font.size = Pt(size - 3); d.font.color.rgb = rgb(dot_color)
            r = p.add_run(); r.text = it; r.font.size = Pt(size); r.font.color.rgb = rgb(INK if it != "—" else MUTED); r.font.name = "Calibri"

    def chip(slide, x, y, rag, w=Inches(1.25)):
        s = box(slide, x, y, w, Inches(0.34), RAG_SOFT[rag], shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        s.adjustments[0] = 0.5
        tf = s.text_frame; tf.margin_left = tf.margin_right = Inches(0.05)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        d = p.add_run(); d.text = "● "; d.font.size = Pt(10); d.font.color.rgb = rgb(RAG_HEX[rag])
        r = p.add_run(); r.text = RAG_LABEL[rag]; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = rgb(RAG_HEX[rag])

    def bar(slide, x, y, w, ratio, color, h=Inches(0.09)):
        box(slide, x, y, w, h, "#E8ECF0", shape=MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.5
        if ratio and ratio > 0:
            b = box(slide, x, y, max(int(h), int(w * ratio)), h, color, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            b.adjustments[0] = 0.5

    def chrome(slide, title, subtitle=None):
        box(slide, 0, 0, SW, Inches(0.12), brand)
        text(slide, Inches(0.5), Inches(0.3), Inches(10), Inches(0.6), title, 26, True)
        if subtitle:
            text(slide, Inches(0.5), Inches(0.88), Inches(11), Inches(0.4), subtitle, 13, False, MUTED)
        text(slide, Inches(0.5), Inches(7.05), Inches(8), Inches(0.3),
             f"{snap.get('org', '')} · {snap.get('title', '')} · Week of {snap['week_label']}", 9, False, MUTED)

    t = snap["totals"]
    projects = snap["projects"]

    # ---- 1. title
    s = prs.slides.add_slide(blank)
    box(s, 0, 0, SW, prs.slide_height, brand)
    box(s, Inches(0.8), Inches(2.55), Inches(0.12), Inches(1.9), "#FFFFFF")
    text(s, Inches(1.15), Inches(2.4), Inches(11), Inches(0.9), snap.get("title", "Weekly Project Progress Report"), 38, True, "#FFFFFF")
    text(s, Inches(1.15), Inches(3.3), Inches(11), Inches(0.6), f"Week of {snap['week_label']}  ·  {scope_label(snap)}", 20, False, "#DCE9EF")
    text(s, Inches(1.15), Inches(3.9), Inches(11), Inches(0.5), snap.get("org", ""), 16, False, "#DCE9EF")
    text(s, Inches(1.15), Inches(6.6), Inches(11), Inches(0.4), f"Generated {snap.get('generated_at', '')}", 11, False, "#9DBCCB")

    # ---- 2. portfolio overview
    s = prs.slides.add_slide(blank)
    chrome(s, "Portfolio at a glance", f"{t['projects']} projects · {t['submitted']} of {t['expected']} updates submitted this week")
    tiles = [("On track", t["G"], RAG_HEX["G"]), ("At risk", t["A"], RAG_HEX["A"]), ("Off track", t["R"], RAG_HEX["R"]),
             ("Open issues", t["open_issues"], INK), ("High / critical", t["critical_issues"], RAG_HEX["R"] if t["critical_issues"] else INK),
             ("Overdue milestones", t["overdue_milestones"], RAG_HEX["A"] if t["overdue_milestones"] else INK)]
    tw, gap = Inches(1.95), Inches(0.12)
    for i, (lbl, val, col) in enumerate(tiles):
        x = Inches(0.5) + i * (tw + gap)
        box(s, x, Inches(1.45), tw, Inches(1.0), "#F7F9FA", LINE)
        text(s, x + Inches(0.12), Inches(1.5), tw, Inches(0.55), str(val), 28, True, col)
        text(s, x + Inches(0.12), Inches(2.05), tw, Inches(0.3), lbl, 11, False, MUTED)

    rows = projects[:10]
    tbl = s.shapes.add_table(len(rows) + 1, 5, Inches(0.5), Inches(2.75), Inches(12.33), Inches(0.4) * (len(rows) + 1)).table
    widths = [3.6, 1.2, 1.4, 1.0, 5.13]
    for i, wdt in enumerate(widths):
        tbl.columns[i].width = Inches(wdt)
    for ci, h in enumerate(["Project", "Phase", "Status", "Progress", "Headline"]):
        c = tbl.cell(0, ci); c.text = h
        c.fill.solid(); c.fill.fore_color.rgb = rgb("#EEF2F5")
        f = c.text_frame.paragraphs[0].runs[0].font; f.size = Pt(11); f.bold = True; f.color.rgb = rgb(MUTED)
    for ri, p in enumerate(rows, start=1):
        vals = [p["name"] + ("  (no update)" if p.get("stale") else ""), p.get("phase") or "",
                "● " + RAG_LABEL[p.get("rag")], f"{p['progress']}%" if p.get("progress") is not None else "—",
                p.get("summary") or "—"]
        for ci, v in enumerate(vals):
            c = tbl.cell(ri, ci); c.text = v
            c.fill.solid(); c.fill.fore_color.rgb = rgb("#FFFFFF")
            c.margin_top = c.margin_bottom = Inches(0.04)
            f = c.text_frame.paragraphs[0].runs[0].font
            f.size = Pt(10.5 if ci == 4 else 11.5); f.color.rgb = rgb(INK)
            if ci == 0: f.bold = True
            if ci == 2: f.bold = True; f.color.rgb = rgb(RAG_HEX[p.get("rag")])
    if len(projects) > 10:
        text(s, Inches(0.5), Inches(6.7), Inches(8), Inches(0.3), f"+{len(projects) - 10} more projects — see detail slides", 10, False, MUTED)

    # ---- 3. one slide per project
    for p in projects:
        s = prs.slides.add_slide(blank)
        sub = f"{p.get('phase') or ''} · Owner: {p.get('owner') or '—'}"
        if p.get("target_date"):
            sub += f" · Target {fmt_date(p['target_date'])}"
        if p.get("trend") in ("up", "down"):
            sub += f" · {TREND[p['trend']]} vs last update"
        chrome(s, p["name"], sub)
        chip(s, Inches(11.55), Inches(0.38), p.get("rag"))
        y = Inches(1.4)
        if p.get("stale"):
            text(s, Inches(0.5), y, Inches(12), Inches(0.3), f"No update this week — showing {fmt_date(p.get('update_week')) if p.get('update_week') else 'no data'}", 11, True, RAG_HEX["A"])
            y += Inches(0.32)
        box(s, Inches(0.5), y, Inches(12.33), Inches(0.7), "#F4F7F9")
        text(s, Inches(0.65), y + Inches(0.05), Inches(12.0), Inches(0.6), p.get("summary") or "No summary provided.", 14, False, INK, anchor=MSO_ANCHOR.MIDDLE)
        top = y + Inches(0.95)

        # left: metrics
        text(s, Inches(0.5), top, Inches(3.8), Inches(0.3), "KEY METRICS", 10, True, MUTED)
        my = top + Inches(0.38)
        for m in (p.get("metrics") or [])[:6]:
            r = metric_ratio(m)
            text(s, Inches(0.5), my, Inches(2.9), Inches(0.3), m["label"], 11.5)
            text(s, Inches(3.2), my, Inches(1.0), Inches(0.3), f"{fmt_num(m['value'])}{m.get('unit') or ''}", 13, True, align=PP_ALIGN.RIGHT)
            if r is not None:
                bar(s, Inches(0.55), my + Inches(0.36), Inches(3.6), r, metric_color(r))
            my += Inches(0.6)
        if not p.get("metrics"):
            text(s, Inches(0.5), my, Inches(3.6), Inches(0.3), "—", 11, False, MUTED); my += Inches(0.4)
        ms = p.get("milestones_overdue", []) + [m for m in p.get("milestones_due", []) if m not in p.get("milestones_overdue", [])]
        if ms and my < Inches(5.6):
            text(s, Inches(0.5), my + Inches(0.1), Inches(3.8), Inches(0.3), "MILESTONES", 10, True, MUTED)
            my += Inches(0.42)
            for m in ms[:3]:
                late = m in p.get("milestones_overdue", [])
                text(s, Inches(0.5), my, Inches(3.8), Inches(0.3),
                     f"{'OVERDUE · ' if late else ''}{m['title']} · {fmt_date(m['due_date'])}", 10.5, late, RAG_HEX["R"] if late else INK)
                my += Inches(0.32)

        # middle + right columns
        colh = Inches(6.9) - top
        bullet_list(s, Inches(4.6), top, Inches(4.1), colh, "Key items in the critical path", p.get("critical_path"), RAG_HEX["R"])
        right_items = p.get("achievements")
        bullet_list(s, Inches(8.9), top, Inches(3.95), colh / 2, "Achieved this week", right_items, RAG_HEX["G"], max_items=4)
        nxt = p.get("next_steps") or []
        if p.get("support_needed"):
            nxt = [f"Support needed: {x}" for x in p["support_needed"]] + nxt
        bullet_list(s, Inches(8.9), top + colh / 2, Inches(3.95), colh / 2, "Next steps", nxt, MUTED, max_items=4)

    # ---- 4. issues
    issues = snap["issues"]
    if issues or snap.get("resolved"):
        for chunk_start in range(0, max(len(issues), 1), 8):
            chunk = issues[chunk_start:chunk_start + 8]
            s = prs.slides.add_slide(blank)
            chrome(s, "Production issues & open risks", f"{t['open_issues']} open · {t['production_issues']} production issues being monitored")
            if not chunk:
                break
            tbl = s.shapes.add_table(len(chunk) + 1, 6, Inches(0.5), Inches(1.45), Inches(12.33), Inches(0.5) * (len(chunk) + 1)).table
            for i, wdt in enumerate([3.4, 2.0, 1.1, 1.3, 1.0, 3.53]):
                tbl.columns[i].width = Inches(wdt)
            for ci, h in enumerate(["Issue", "Project", "Severity", "Status", "Due", "Latest progress"]):
                c = tbl.cell(0, ci); c.text = h
                c.fill.solid(); c.fill.fore_color.rgb = rgb("#EEF2F5")
                f = c.text_frame.paragraphs[0].runs[0].font; f.size = Pt(11); f.bold = True; f.color.rgb = rgb(MUTED)
            for ri, i in enumerate(chunk, start=1):
                due = fmt_date(i.get("due_date")) + (" (overdue)" if i.get("overdue") else "")
                vals = [f"{i['title']}\n{i['kind']}", i.get("project") or "—", i["severity"], i["status"], due, i.get("latest_note") or "—"]
                for ci, v in enumerate(vals):
                    c = tbl.cell(ri, ci); c.text = v
                    c.fill.solid(); c.fill.fore_color.rgb = rgb("#FFFFFF")
                    for pi, para in enumerate(c.text_frame.paragraphs):
                        for r in para.runs:
                            r.font.size = Pt(9.5 if pi else 11); r.font.color.rgb = rgb(MUTED if pi else INK)
                            if ci == 0 and pi == 0: r.font.bold = True
                            if ci == 2: r.font.bold = True; r.font.color.rgb = rgb(SEV_HEX.get(i["severity"], MUTED))
                            if ci == 4 and i.get("overdue"): r.font.color.rgb = rgb(RAG_HEX["R"]); r.font.bold = True
            if snap.get("resolved") and chunk_start + 8 >= len(issues):
                text(s, Inches(0.5), Inches(6.6), Inches(12), Inches(0.4),
                     "Resolved this week: " + "; ".join(i["title"] for i in snap["resolved"]), 11, False, RAG_HEX["G"])

    prs.core_properties.title = f"{snap.get('title')} — {snap['week_label']}"
    prs.save(out)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--input", required=True)
    ap.add_argument("--output", required=True)
    ap.add_argument("--format", choices=["pdf", "pptx"], required=True)
    a = ap.parse_args(argv)
    with open(a.input, encoding="utf-8") as fh:
        snap = json.load(fh)
    (render_pdf if a.format == "pdf" else render_pptx)(snap, a.output)
    print(a.output)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:  # surface a clean message to Laravel's Process output
        print(f"report error: {exc}", file=sys.stderr)
        raise
